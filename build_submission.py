import argparse
import ast
import hashlib
import html
import json
import math
import os
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from PIL import Image
from playwright.sync_api import sync_playwright


ROOT = Path(__file__).resolve().parent

if Path(sys.prefix).resolve() != (ROOT / ".venv").resolve():
    raise SystemExit("Run this report generator using .venv/bin/python.")


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def paragraph(text, style=""):
    return f'<p class="{style}">{html.escape(str(text))}</p>'


def table(frame, title=None, index=False):
    heading = f"<h2>{html.escape(title)}</h2>" if title else ""
    return heading + frame.to_html(index=index, border=0, classes="data", float_format=lambda value: f"{value:.4f}", escape=True)


def console(text):
    return '<pre class="console">' + html.escape(text.strip()) + "</pre>"


def log_section(log, start, end=None):
    position = log.find(start)
    if position < 0:
        raise ValueError(f"Execution log does not contain {start!r}.")
    stop = log.find(end, position + len(start)) if end else len(log)
    return log[position:stop if stop >= 0 else len(log)].strip()


def make_pages(source, frame, output, config, metrics, log):
    pages = []
    lines = source.splitlines()
    nodes = {node.name: node for node in ast.parse(source).body if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))}
    cleaned = pd.read_csv(output / "cleaned_dataset.csv")
    text_column = config["text_column"]
    label_column = config["label_column"]
    label_count = len(config["classes"])
    split_counts = config["split_counts"]
    cleaned_count = len(cleaned)
    limitations = (
        f"The supplied file contains {len(frame):,} rows but only {frame[text_column].nunique():,} distinct raw comments. "
        f"After cleaning and removal of repeated comments, {cleaned_count} unique examples remain across {label_count} classes. "
        f"Training, validation and test contain {split_counts['train']}, {split_counts['validation']} and {split_counts['test']} comments respectively. "
        f"One test prediction changes accuracy by {100 / split_counts['test']:.0f} percentage points. "
        "Identical cleaned text is kept in one partition only. These results demonstrate the workflow; the very small independent sample cannot establish dependable real-world accuracy."
    )

    def add(task, title, content, sources, note=None):
        pages.append({"task": task, "title": title, "content": content, "sources": sources, "note": note or ""})

    def code(task, title, names):
        selected = []
        for name in names:
            if name not in nodes:
                raise ValueError(f"Required function is absent from main.py: {name}")
            node = nodes[name]
            selected.extend((number + 1, lines[number]) for number in range(node.lineno - 1, node.end_lineno))
            selected.append((None, ""))
        batches = []
        batch = []
        cost = 0
        for number, line in selected:
            line_cost = max(1, math.ceil(len(line.expandtabs(4)) / 86))
            if batch and cost + line_cost > 45:
                batches.append(batch)
                batch = []
                cost = 0
            batch.append((number, line))
            cost += line_cost
        if batch:
            batches.append(batch)
        for number, batch in enumerate(batches, start=1):
            cells = "".join(
                '<div class="source-row"><span class="line-number">'
                + (str(line_number) if line_number is not None else "")
                + '</span><code>' + html.escape(line) + "</code></div>"
                for line_number, line in batch
            )
            suffix = f" ({number}/{len(batches)})" if len(batches) > 1 else ""
            add(task, title + suffix, '<div class="source">' + cells + "</div>", ["main.py"], "Exact source from the executed main.py; left-hand numbers identify original source lines.")

    code("Task 1", "Dataset loading and exploration: Python code", ["task_one_load_and_explore"])
    selected_columns = [text_column, label_column]
    id_columns = [column for column in frame.columns if "id" in column.lower() and column not in selected_columns]
    first_band = id_columns[:1] + selected_columns
    remaining = [column for column in frame.columns if column not in first_band]
    numeric_columns = [column for column in remaining if pd.api.types.is_numeric_dtype(frame[column])]
    other_columns = [column for column in remaining if column not in numeric_columns]
    bands = [first_band]
    if numeric_columns:
        bands.extend([id_columns[:1] + numeric_columns[offset:offset + 9] for offset in range(0, len(numeric_columns), 9)])
    if other_columns:
        bands.append(id_columns[:1] + other_columns)
    for number, columns in enumerate(bands, start=1):
        content = table(frame.head(10)[columns], "First 10 records") + table(frame.tail(8)[columns], "Last 8 records")
        add("Task 1", f"First and last records: column band {number}/{len(bands)}", content, [str(Path(config["data"]).name), "outputs/run.log"], "The same required records are shown in column bands so every original column remains readable.")
    dtypes = pd.DataFrame({"Column": frame.columns, "Data type": [str(value) for value in frame.dtypes]})
    add("Task 1", "Dimensions, column names and data types", paragraph(f"Dataset shape: {frame.shape[0]:,} rows × {frame.shape[1]} columns.", "highlight") + table(dtypes), [str(Path(config["data"]).name), "outputs/run.log"])
    numerical = frame.select_dtypes(include="number")
    if numerical.empty:
        add("Task 1", "Numerical statistical summary", paragraph("No numerical columns are present."), ["outputs/run.log"])
    else:
        summary = numerical.describe()
        for offset in range(0, len(summary.columns), 5):
            band = summary.iloc[:, offset:offset + 5]
            add("Task 1", f"Numerical statistical summary ({offset // 5 + 1})", table(band, index=True), [str(Path(config["data"]).name), "outputs/run.log"], "Pandas describe() reports count, mean, standard deviation, minimum, quartiles and maximum.")

    code("Task 2", "Cleaning and data quality: Python code", ["clean_text", "task_two_prepare"])
    missing = pd.read_csv(output / "missing_values.csv")
    exact_duplicates = int(frame.duplicated().sum())
    duplicate_comments = int(frame[text_column].duplicated().sum())
    data_quality = pd.DataFrame({"Check": ["Raw rows", "Exact duplicate records beyond first", "Repeated raw comments beyond first", "Unique raw comments", "Cleaned unique comments"], "Value": [len(frame), exact_duplicates, duplicate_comments, frame[text_column].nunique(), cleaned_count]})
    add("Task 2", "Missing values and duplicates", table(data_quality) + table(missing, "Missing values per column"), ["outputs/missing_values.csv", str(Path(config["data"]).name), "outputs/run.log"])
    raw_counts = frame[label_column].value_counts().sort_index().rename("Raw rows")
    clean_counts = cleaned["label"].value_counts().sort_index().rename("Unique cleaned comments")
    distribution = pd.concat([raw_counts, clean_counts], axis=1).fillna(0).astype(int)
    add("Task 2", "Class distribution and dataset limitation", table(distribution, index=True) + paragraph(limitations, "limitation"), ["outputs/cleaned_dataset.csv", "outputs/class_distribution.csv", str(Path(config["data"]).name)])
    examples = cleaned[["raw_text", "cleaned_text", "label"]].head(10)
    add("Task 2", "Cleaning examples", table(examples), ["outputs/cleaned_dataset.csv"], "Lowercase text retains alphabetic characters and spaces; tags, URLs, punctuation, digits, emoji and excess whitespace are removed.")

    split_helpers = [name for name in nodes if "split" in name and name != "task_three_tokenize"]
    code("Task 3", "Splitting, tokenization and padding: Python code", split_helpers + ["task_three_tokenize"])
    third_log = log_section(log, "TASK 3:", "TASK 4:")
    third_lines = third_log.splitlines()
    for offset in range(0, len(third_lines), 30):
        add("Task 3", "Tokenizer and padded sequence output" + (f" ({offset // 30 + 1})" if len(third_lines) > 30 else ""), console("\n".join(third_lines[offset:offset + 30])), ["outputs/run.log", "outputs/tokenizer.json"], "Tokenizer vocabulary is fitted only on training comments. Unseen words use the out-of-vocabulary token; padding and truncation are applied after splitting.")

    code("Task 4", "LSTM architecture and training: Python code", ["make_dataset", "task_four_train"])
    fourth_log = log_section(log, "TASK 4:", "TASK 5:")
    model_part = fourth_log.split("Epoch 1/", 1)[0].strip()
    if len(model_part.splitlines()) > 35:
        model_part = "\n".join(model_part.splitlines()[:35])
    command = f".venv/bin/python main.py --data {Path(config['data']).name} --epochs {config['epochs']} --max-length {config['max_length']} --batch-size {config['batch_size']} --seed {config['seed']}"
    add("Task 4", "Model architecture and training configuration", console(model_part) + paragraph(f"Requested epochs: {config['epochs']}; batch size: {config['batch_size']}; maximum sequence length: {config['max_length']}; vocabulary size: {config['vocabulary_size']}.") + paragraph("Command to reproduce these settings:") + console(command), ["outputs/run.log", "outputs/run_config.json"], "Embedding → LSTM → Dropout → Dense with softmax. Integer targets use sparse categorical cross-entropy. The checkpoint is selected using validation loss.")
    history = pd.read_csv(output / "training_history.csv")
    add("Task 4", "Training and validation results for every epoch", table(history) + paragraph(f"Best validation-loss checkpoint: epoch {metrics['best_epoch']}.", "highlight"), ["outputs/training_history.csv", "outputs/metrics.json"], "All configured epochs were executed. Test examples were excluded from model fitting and checkpoint selection.")

    code("Task 5", "Held-out evaluation: Python code", ["task_five_evaluate"])
    result_text = f"Test accuracy: {metrics['test_accuracy']:.4f} ({metrics['test_accuracy'] * 100:.1f}%)\nTest loss: {metrics['test_loss']:.4f}\nMacro F1-score: {metrics['macro_f1']:.4f}\nWeighted F1-score: {metrics['weighted_f1']:.4f}\nTest examples: {split_counts['test']}"
    report = (output / "classification_report.txt").read_text(encoding="utf-8")
    add("Task 5", "Test metrics and classification report", console(result_text) + console(report) + paragraph(limitations, "limitation"), ["outputs/metrics.json", "outputs/classification_report.txt"])
    for filename, title, note in [
        ("confusion_matrix.png", "Confusion matrix on the test set", "Rows show actual classes; columns show predicted classes. Counts come from the held-out test set."),
        ("training_curves.png", "Training and validation accuracy/loss curves", "Curves report all training epochs. Training loss uses class weights; validation loss is unweighted. Accuracy changes in large steps because each partition contains very few distinct comments."),
    ]:
        add("Task 5", title, '<img class="figure" src="' + (output / filename).resolve().as_uri() + '" alt="' + html.escape(title) + '">', ["outputs/" + filename], note)
    test_predictions = pd.read_csv(output / "test_predictions.csv")
    visible_test = test_predictions[["actual", "predicted"]].copy()
    rows = pd.read_csv(output / "test_rows.csv")
    if "source_row" in rows.columns and len(rows) == len(visible_test):
        visible_test.insert(0, "CSV row", rows["source_row"].tolist())
    add("Task 5", "Individual held-out predictions", table(visible_test) + paragraph("Each row is one independent test comment. Prediction probabilities and full metrics remain in the saved output files."), ["outputs/test_predictions.csv", "outputs/test_rows.csv"])
    code("Task 5", "Predictions for new comments: Python code", ["predict_new_comments"])
    new_predictions = pd.read_csv(output / "new_comment_predictions.csv")
    compact_columns = [column for column in new_predictions.columns if not column.startswith("probability_") and column != "cleaned_text"]
    add("Task 5", "New-comment prediction output", table(new_predictions[compact_columns]) + paragraph("These new comments have no supplied ground-truth labels. Predictions demonstrate inference only and are excluded from the reported test accuracy.", "limitation"), ["outputs/new_comment_predictions.csv"])
    return pages


def html_page(item, number, total):
    css = """
    * { box-sizing: border-box; }
    body { margin: 0; background: white; color: rgb(26, 37, 55); font-family: Arial, sans-serif; }
    article { width: 1200px; padding: 34px 38px 26px; background: white; }
    .eyebrow { color: rgb(32, 102, 138); font-size: 18px; font-weight: 700; letter-spacing: 1px; margin-bottom: 12px; }
    h1 { font-size: 32px; line-height: 1.2; margin: 0 0 14px; }
    h2 { font-size: 23px; margin: 22px 0 10px; color: rgb(32, 102, 138); }
    p { font-size: 22px; line-height: 1.42; margin: 16px 0; }
    .note { font-size: 18px; line-height: 1.4; color: rgb(72, 86, 104); padding-bottom: 12px; }
    .highlight { background: rgb(232, 243, 249); padding: 16px; font-weight: 700; }
    .limitation { padding: 18px; background: rgb(255, 247, 221); border-left: 5px solid rgb(197, 149, 32); }
    table { width: 100%; border-collapse: collapse; table-layout: auto; margin: 12px 0 24px; }
    th { background: rgb(232, 239, 246); text-align: left; font-size: 19px; padding: 10px 8px; overflow-wrap: anywhere; }
    td { font-size: 20px; padding: 9px 8px; border-bottom: 1px solid rgb(213, 222, 231); vertical-align: top; overflow-wrap: anywhere; }
    tbody tr:nth-child(even) { background: rgb(248, 250, 253); }
    .source { background: rgb(246, 248, 251); border: 1px solid rgb(217, 225, 233); padding: 12px 8px; }
    .source-row { display: flex; min-height: 26px; }
    .line-number { color: rgb(120, 132, 147); width: 53px; flex: 0 0 53px; font: 17px/26.25px monospace; text-align: right; padding-right: 12px; user-select: none; }
    code { white-space: pre-wrap; overflow-wrap: anywhere; word-break: break-all; min-width: 0; font: 21px/26.25px 'DejaVu Sans Mono', monospace; }
    .console { white-space: pre-wrap; overflow-wrap: anywhere; word-break: break-all; font: 20px/1.4 'DejaVu Sans Mono', monospace; padding: 18px; border: 1px solid rgb(217, 225, 233); background: rgb(246, 248, 251); }
    .figure { width: 100%; height: auto; max-height: 1150px; object-fit: contain; }
    footer { border-top: 1px solid rgb(213, 222, 231); margin-top: 24px; padding-top: 12px; color: rgb(90, 105, 122); font-size: 16px; line-height: 1.4; }
    """
    heading = '<div class="eyebrow">TOXIC COMMENT CLASSIFICATION · ' + html.escape(item["task"].upper()) + "</div>"
    title = "<h1>" + html.escape(item["title"]) + "</h1>"
    note = paragraph(item["note"], "note") if item["note"] else ""
    footer = "<footer>Exported execution evidence · " + html.escape(", ".join(item["sources"])) + f"<br>Screenshot {number} of {total} · Generated from the executed source and saved run artifacts</footer>"
    return '<!DOCTYPE html><html lang="en"><head><meta charset="utf-8"><title>' + html.escape(item["title"]) + "</title><style>" + css + "</style></head><body><article>" + heading + title + note + item["content"] + footer + "</article></body></html>"


def validate_run(output, source_path, config):
    expected = ["cleaned_dataset.csv", "missing_values.csv", "class_distribution.csv", "train_rows.csv", "validation_rows.csv", "test_rows.csv", "tokenizer.json", "classes.json", "training_history.csv", "training_curves.png", "best_model.keras", "confusion_matrix.csv", "confusion_matrix.png", "classification_report.txt", "test_predictions.csv", "metrics.json", "run_config.json", "new_comment_predictions.csv", "run.log"]
    missing = [name for name in expected if not (output / name).is_file()]
    if missing:
        raise FileNotFoundError("Run the complete assignment first. Missing output files: " + ", ".join(missing))
    data_path = Path(config["data"])
    if not data_path.is_file():
        raise FileNotFoundError(f"The executed dataset is missing: {data_path}")
    for path, key in [(source_path, "main_script_sha256"), (data_path, "dataset_sha256")]:
        if key in config and digest(path) != config[key]:
            raise ValueError(f"The current {path.name} does not match the file used in the saved run.")
        if key not in config and path.stat().st_mtime > (output / "run_config.json").stat().st_mtime:
            raise ValueError(f"The current {path.name} is newer than the completed run. Execute main.py again.")
    source = source_path.read_text(encoding="utf-8")
    if chr(35) in source:
        raise ValueError("main.py contains a forbidden hash character.")
    history = pd.read_csv(output / "training_history.csv")
    if len(history) != config["epochs"] or not 10 <= len(history) <= 20:
        raise ValueError("Saved training history does not contain the configured 10–20 epochs.")
    partitions = {name: pd.read_csv(output / f"{name}_rows.csv") for name in ("train", "validation", "test")}
    for name, partition in partitions.items():
        if len(partition) != config["split_counts"][name]:
            raise ValueError(f"The saved {name} partition disagrees with run_config.json.")
    source_rows = [set(partition["source_row"]) for partition in partitions.values()]
    if any(source_rows[left] & source_rows[right] for left in range(3) for right in range(left + 1, 3)):
        raise ValueError("Saved partitions overlap.")
    predictions = pd.read_csv(output / "test_predictions.csv")
    metrics = read_json(output / "metrics.json")
    actual_accuracy = float(predictions["actual"].eq(predictions["predicted"]).mean())
    if len(predictions) != config["split_counts"]["test"] or not math.isclose(actual_accuracy, metrics["test_accuracy"], abs_tol=1e-6):
        raise ValueError("Saved predictions and test metrics are inconsistent.")
    matrix = pd.read_csv(output / "confusion_matrix.csv", index_col=0)
    if int(matrix.to_numpy().sum()) != len(predictions):
        raise ValueError("The confusion matrix has a different test sample count.")
    return data_path, source, metrics


def main():
    parser = argparse.ArgumentParser(description="Create the screenshot-only Word assignment from a completed real-data run.")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "outputs")
    parser.add_argument("--submission-dir", type=Path, default=ROOT / "submission")
    parser.add_argument("--chrome", type=Path, default=Path("/usr/bin/google-chrome"))
    arguments = parser.parse_args()
    output = arguments.output_dir.resolve()
    destination = arguments.submission_dir.resolve()
    if not destination.is_relative_to(ROOT):
        raise ValueError("The submission directory must remain inside the project folder.")
    config = read_json(output / "run_config.json")
    source_path = ROOT / "main.py"
    data_path, source, metrics = validate_run(output, source_path, config)
    frame = pd.read_csv(data_path)
    raw_log = (output / "run.log").read_text(encoding="utf-8")
    log = re.sub(r"\x1b\[[0-9;]*[A-Za-z]", "", raw_log)
    if "All five tasks completed" not in log:
        raise ValueError("The captured run log does not confirm successful completion.")
    pages = make_pages(source, frame, output, config, metrics, log)
    screenshot_dir = destination / "screenshots"
    page_dir = destination / "pages"
    screenshot_dir.mkdir(parents=True, exist_ok=True)
    page_dir.mkdir(parents=True, exist_ok=True)
    manifest = {"generated_at_utc": datetime.now(timezone.utc).isoformat(), "main_script_sha256": digest(source_path), "dataset_sha256": digest(data_path), "run_config_sha256": digest(output / "run_config.json"), "method": "Actual Chrome screenshots of local HTML containing exact executed source, console excerpts, and saved output artifacts.", "pages": []}
    screenshots = []
    for variable, directory in {"TMPDIR": "tmp", "XDG_CACHE_HOME": "xdg", "XDG_CONFIG_HOME": "config"}.items():
        location = ROOT / ".cache" / "report" / directory
        location.mkdir(parents=True, exist_ok=True)
        os.environ[variable] = str(location)
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(executable_path=str(arguments.chrome), headless=True, args=["--disable-dev-shm-usage"])
        page = browser.new_page(viewport={"width": 1200, "height": 1550}, device_scale_factor=1.5)
        for index, item in enumerate(pages, start=1):
            page_path = page_dir / f"{index:02d}.html"
            page_path.write_text(html_page(item, index, len(pages)), encoding="utf-8")
            page.goto(page_path.as_uri(), wait_until="networkidle")
            page.evaluate("document.fonts.ready")
            page.locator("article img").evaluate_all("images => Promise.all(images.map(image => image.decode()))")
            width = page.evaluate("Math.max(document.documentElement.scrollWidth, document.body.scrollWidth)")
            bounds = page.locator("article").bounding_box()
            if width > 1200 or bounds["height"] > 1750:
                raise ValueError(f"Screenshot {index} needs pagination: width={width}, height={bounds['height']}")
            screenshot_path = screenshot_dir / f"{index:02d}.png"
            page.locator("article").screenshot(path=str(screenshot_path))
            bounds = page.locator("article").bounding_box()
            screenshots.append(screenshot_path)
            manifest["pages"].append({"number": index, "task": item["task"], "title": item["title"], "sources": item["sources"], "html": str(page_path.relative_to(destination)), "screenshot": str(screenshot_path.relative_to(destination)), "css_height": bounds["height"], "scroll_width": width, "screenshot_sha256": digest(screenshot_path)})
            print(f"Captured {index}/{len(pages)}: {item['title']} ({int(bounds['height'])} px high)", flush=True)
        browser.close()
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    style = document.styles["Normal"]
    style.font.size = Pt(1)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.line_spacing = 1
    for index, screenshot in enumerate(screenshots):
        if index:
            document.add_page_break()
        with Image.open(screenshot) as picture:
            ratio = picture.height / picture.width
        width = min(7.6, 9.9 / ratio)
        paragraph_object = document.add_paragraph()
        paragraph_object.paragraph_format.keep_together = True
        paragraph_object.add_run().add_picture(str(screenshot), width=Inches(width))
    document.core_properties.title = "Toxic Comment Classification — Tasks 1–5"
    document.core_properties.subject = "Screenshots of executed Python code and genuine real-dataset outputs"
    document.core_properties.author = ""
    document.core_properties.comments = "Generated from the supplied CSV and completed execution; screenshots are not simulated IDE windows."
    document_path = destination / "Toxic_Comment_Classification_Assignment.docx"
    document.save(document_path)
    manifest["document"] = document_path.name
    manifest["document_sha256"] = digest(document_path)
    (destination / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    links = "".join('<li><a href="' + item["html"] + '">' + html.escape(str(item["number"]) + ". " + item["task"] + ": " + item["title"]) + "</a></li>" for item in manifest["pages"])
    (destination / "evidence.html").write_text('<!DOCTYPE html><html lang="en"><meta charset="utf-8"><title>Assignment execution evidence</title><body><h1>Toxic Comment Classification</h1><p>Actual browser screenshots of executed source and saved output artifacts.</p><ol>' + links + "</ol></body></html>", encoding="utf-8")
    print(f"Saved {document_path} with {len(screenshots)} evidence screenshots.")


if __name__ == "__main__":
    main()
